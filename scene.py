# -*- coding: utf-8 -*-
"""Реестр сцены: что из живого сохранения принадлежит проекту.

sync.py по этому списку находит объекты в файле и обновляет их, когда
меняются исходники. Здесь же — функции, которые из строки таблицы делают
текст тултипа; они общие для сборки с нуля (gen_save.py), рендера листов
(gen_cards.py) и синхронизации (sync.py).

КОЛОДЫ НАХОДЯТСЯ САМИ. Любая папка `Карты\\<Название>\\cards.xlsx` — колода.
Тип определяется по колонкам таблицы:
    t_name, b_name, affil ...   карта действия, две половины (как у игрока)
    name, effect, tipe          карта снаряжения
Название колоды — имя папки. Чтобы завести новую колоду, достаточно создать
папку с cards.xlsx нужной схемы и запустить sync.py — ничего править в коде
не нужно. Три исходные колоды перечислены в KNOWN только ради того, чтобы
их ключи, раскладка листа и место на столе не изменились.

КАК ОБЪЕКТ ОПОЗНАЁТСЯ В ФАЙЛЕ. Ни GUID, ни номер колоды не годятся: TTS
перенумеровывает колоды при каждом сохранении, а у колод игрока к тому же
пропали названия. Поэтому sync.py метит каждый свой объект в поле GMNotes
строкой вида `teyrin:map_island@d7ba6b15` — ключ ассета и хеш содержимого.
По ключу объект находится, по хешу видно, устарел ли. Пока метки нет,
объект опознаётся по названию (доски, фигурки) или по размеру листа (колоды).
"""
import math, os, re
import openpyxl

from paths import CARDS

MARK = "teyrin:"      # префикс метки в GMNotes


# ------------------------------------------------------------------ тултипы
def act_card(c):
    """Карта действия: две половины, принадлежность."""
    nick = c["t_name"] + (" / " + c["b_name"] if c["b_name"] else "")
    parts = ["Верх: {} — ОД {}, вын. {}\n{}".format(c["t_name"], c["t_ap"] or "—",
                                                    c["t_stam"] or "—", c["t_effect"])]
    if c["b_name"]:
        parts.append("Низ: {} — ОД {}, вын. {}\n{}".format(c["b_name"], c["b_ap"] or "—",
                                                           c["b_stam"] or "—", c["b_effect"]))
    if c["affil"] and c["affil"] != "—":
        parts.append("Принадлежность: " + c["affil"])
    return nick, "\n\n".join(parts)


def gear_card(c):
    return c["name"], "{}\n\nТип: {}".format(c["effect"], c["tipe"])


KINDS = {
    # kind     колонка-признак  колонка имени  тултип     уникальные рубашки
    "action": dict(marker="t_name", name_col="t_name", describe=act_card, unique_back=True),
    "gear":   dict(marker="effect", name_col="name",   describe=gear_card, unique_back=False),
}


# ----------------------------------------------------------- чтение таблиц
def read_cards(xlsx, name_col):
    ws = openpyxl.load_workbook(xlsx, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    out = []
    for r in rows[1:]:
        c = {h: ("" if v is None else str(v).strip()) for h, v in zip(hdr, r)}
        if c.get(name_col):                  # пустая строка-заготовка пропускается
            out.append(c)
    return out


def header(xlsx):
    ws = openpyxl.load_workbook(xlsx, read_only=True).active
    return [h for h in next(ws.iter_rows(values_only=True)) if h]


# ------------------------------------------------------------------- колоды
# Исходные три колоды: ключ, раскладка листа и место фиксированы, чтобы
# ссылки в сцене не менялись только из-за перестановки.
KNOWN = {
    "Карты действий":   dict(key="actions", sheet=(5, 4), pos=(-8.0, -10.5)),
    "Карты снаряжения": dict(key="gear",    sheet=(3, 1), pos=(-3.5, -10.5)),
    "Карты зверя":      dict(key="beast",   sheet=(4, 2), pos=(12.0, 11.0)),
}
NEW_DECK_COLUMN = 16.0     # новые колоды встают столбиком у правого края
NEW_DECK_TOP = 11.0
NEW_DECK_STEP = 4.0

_TRANSLIT = dict(zip("абвгдеёжзийклмнопрстуфхцчшщъыьэюя",
                     ["a", "b", "v", "g", "d", "e", "yo", "zh", "z", "i", "y", "k", "l", "m",
                      "n", "o", "p", "r", "s", "t", "u", "f", "h", "c", "ch", "sh", "sch", "",
                      "y", "", "e", "yu", "ya"]))


def slug(name):
    """Имя папки -> ключ ассета: латиница, цифры, подчёркивания."""
    s = "".join(_TRANSLIT.get(ch, ch) for ch in name.lower())
    return re.sub(r"[^a-z0-9]+", "_", s).strip("_") or "deck"


def auto_sheet(n):
    """Раскладка листа под n карт: не шире 10 и не выше 7 — предел TTS."""
    cols = min(n, 10)
    rows = math.ceil(n / cols)
    if rows > 7:
        raise SystemExit("в колоде %d карт, а на один лист влезает 70" % n)
    return cols, rows


def discover_decks():
    """Все колоды из папки Карты: известные с фиксированными настройками,
    остальные — с ключом из имени, листом по числу карт и местом по порядку."""
    out, fresh = [], 0
    for folder in sorted(os.listdir(CARDS)):
        xlsx = os.path.join(CARDS, folder, "cards.xlsx")
        if not os.path.exists(xlsx):
            continue
        hdr = header(xlsx)
        kind = next((k for k, spec in KINDS.items() if spec["marker"] in hdr), None)
        if kind is None:
            print("scene: %s — в cards.xlsx нет ни t_name, ни effect, колода пропущена" % folder)
            continue
        spec = KINDS[kind]
        n = len(read_cards(xlsx, spec["name_col"]))
        if n == 0:
            print("scene: %s — в таблице нет карт, колода пропущена" % folder)
            continue
        known = KNOWN.get(folder, {})
        if "pos" not in known:
            pos = (NEW_DECK_COLUMN, NEW_DECK_TOP - NEW_DECK_STEP * fresh)
            fresh += 1
        else:
            pos = known["pos"]
        out.append(dict(key=known.get("key", slug(folder)), nick=folder, xlsx=xlsx,
                        kind=kind, name_col=spec["name_col"], describe=spec["describe"],
                        unique_back=spec["unique_back"], sheet=known.get("sheet", auto_sheet(n)),
                        pos=pos, count=n))
    return out


DECKS = discover_decks()

# ----------------------------------------------------------------- картинки
# nick — название объекта в сцене, по нему первичное опознание
# kind — Custom_Tile / Figurine_Custom / Custom_Token: как создать, если нет
# pos  — куда положить, если объекта в сцене нет
IMAGES = [
    dict(key="map_island",    nick="Локация 2 — остров",       kind="Custom_Tile",     pos=(0.0, -7.8)),
    dict(key="map_camp",      nick="Локация 1 — старт",        kind="Custom_Tile",     pos=(0.0, 9.5)),
    dict(key="transition",    nick="Переход между локациями",  kind="Custom_Tile",     pos=(0.0, 2.6)),
    dict(key="fig_crab",      nick="Кокосовый краб",           kind="Figurine_Custom", pos=(4.5, -11.7)),
    dict(key="fig_palm",      nick="Пальма с кокосами",        kind="Figurine_Custom", pos=(4.5, -3.0)),
    dict(key="fig_beast",     nick="Зверь из другого мира",    kind="Figurine_Custom", pos=(-3.0, -14.3)),
    dict(key="tok_firestone", nick="Огненный камень",          kind="Custom_Token",    pos=(-4.5, -8.2)),
    dict(key="fig_npc",       nick="NPC",                      kind="Figurine_Custom", pos=(-1.5, 12.6)),
    dict(key="fig_player",    nick="Вы",                       kind="Figurine_Custom", pos=(3.0, 11.7)),
]


if __name__ == "__main__":
    for d in DECKS:
        print("  %-20s %-8s %-7s карт %-3d лист %dx%d  место %s"
              % (d["nick"], d["key"], d["kind"], d["count"], *d["sheet"], d["pos"]))
